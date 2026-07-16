"""
Cross-validates the teacher-sheet-derived class schedule against the
subject-level Class 6/7/8/9&10 sheets. Run this after build_data_v3.py to
sanity-check the result — it's what caught the Shazia/8EE slot mismatch and
is worth re-running any time the source workbook changes meaningfully.

Usage: python3 verify_against_class_sheets.py <file.xlsx>
(reads school_data_v3.json produced by build_data_v3.py in the same directory)
"""
import openpyxl, json, re, sys

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else 'Working_Timetables_2026_-_27_-_Copy.xlsx'
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
school = json.load(open('school_data_v3.json'))
CLASS_SCHEDULE = school['CLASS_SCHEDULE']
SCHOOL_DAYS = ['Monday','Tuesday','Wednesday','Thursday','Friday']

def parse_class_sheet(sheetname):
    ws = wb[sheetname]
    merged = list(ws.merged_cells.ranges)
    def gv(r,c):
        for mr in merged:
            if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
                return ws.cell(row=mr.min_row, column=mr.min_col).value
        return ws.cell(row=r, column=c).value
    results = {}
    max_r, max_c = ws.max_row, ws.max_column
    class_row_starts = []
    for r in range(1, max_r+1):
        c = 1
        while c <= max_c:
            v = gv(r,c)
            if v and re.match(r'^\d{1,2}E[A-E]$', str(v).strip()):
                cls = str(v).strip(); start_c = c; run = 0
                while c <= max_c and gv(r,c) == v:
                    run += 1; c += 1
                if run >= 4:
                    class_row_starts.append((r, start_c, cls))
                continue
            c += 1
    class_row_starts.sort()
    clusters = {}
    for (r, startcol, cls) in class_row_starts:
        clusters.setdefault(startcol, []).append((r, cls))
    DAYSET = {'Monday','Tuesday','Wed','Thursday','Friday'}
    DAYNORM = {'Monday':'Monday','Tuesday':'Tuesday','Wed':'Wednesday','Thursday':'Thursday','Friday':'Friday'}
    for startcol, items in clusters.items():
        items.sort()
        for i, (r, cls) in enumerate(items):
            end_r = items[i+1][0] if i+1 < len(items) else max_r+1
            for rr in range(r, end_r):
                dcol = None
                for cand in (startcol-1, startcol, startcol-2):
                    if cand < 1: continue
                    v = gv(rr, cand)
                    if v and str(v).strip() in DAYSET:
                        dcol = cand; break
                if dcol is None: continue
                day = DAYNORM[str(gv(rr, dcol)).strip()]
                slots = [str(gv(rr,cc)).strip() if gv(rr,cc) else '' for cc in range(dcol+1, dcol+9)]
                results.setdefault(cls, {})
                if day in results[cls]:
                    existing = results[cls][day]
                    results[cls][day] = [e if e else s for e,s in zip(existing, slots)]
                else:
                    results[cls][day] = slots
    return results

subject_data = {}
for s in ['Class 6','Class 7','Class 8','9 & 10']:
    subject_data.update(parse_class_sheet(s))

def norm_subject_label(s):
    s = s.upper().replace('.', '').strip()
    mapping = {
        'MATH':'MATHS','MATHS':'MATHS', 'URDU':'URD','URD':'URD',
        'CHE':'SCI','CHEM':'SCI','PHY':'SCI','BIO':'SCI','SCI':'SCI','CS/BIO':'SCI','BIO/CS':'SCI',
        'PST':'SST','SST':'SST', 'CS':'ICT','ICT':'ICT',
        'ISL':'ISLSIND','SIND':'ISLSIND','SINDH':'ISLSIND',
        'PE':'PE','ENG':'ENG','LIB':'LIB',
    }
    return mapping.get(s, s)

EXTRA_SUBJECT = {'Sher Zaman': 'PE', 'Fazila': 'SCI', 'Nizar': 'MATHS', 'Saleem - NT': 'SST'}

def infer_subject(teacher):
    if teacher in EXTRA_SUBJECT: return EXTRA_SUBJECT[teacher]
    t = teacher.lower()
    if 'eng' in t: return 'ENG'
    if 'sst' in t or 'pst' in t: return 'SST'
    if 'bio' in t or 'chem' in t or 'phys' in t or 'sci' in t: return 'SCI'
    if 'math' in t: return 'MATHS'
    if 'urdu' in t: return 'URD'
    if 'islamiat' in t or 'sindhi' in t or re.search(r'\bisl\b', t): return 'ISLSIND'
    if 'ict' in t or 'computer' in t: return 'ICT'
    if 'lib' in t: return 'LIB'
    return None

mismatches, gaps, extra = [], [], []
total_checked = total_match = 0
for cls in sorted(subject_data.keys()):
    for day in SCHOOL_DAYS:
        subj_row = subject_data.get(cls, {}).get(day, ['']*8)
        entries = {e['slot']: e for e in CLASS_SCHEDULE.get(cls, {}).get(day, [])}
        for idx, subj in enumerate(subj_row):
            slot = idx + 1
            subj = subj.strip().upper()
            entry = entries.get(slot)
            if subj and subj not in ('REM',''):
                canon = norm_subject_label(subj)
                if canon == 'LIB': continue
                total_checked += 1
                if not entry:
                    gaps.append((cls, day, slot, subj)); continue
                inferred = infer_subject(entry['teacher'])
                if inferred == canon: total_match += 1
                else: mismatches.append((cls, day, slot, subj, canon, entry['teacher'], inferred))
            elif subj == '' and entry:
                extra.append((cls, day, slot, entry['teacher']))

print(f"Checked {total_checked} cells against the Class sheets. Clean matches: {total_match}.")
print(f"\nGenuine subject mismatches ({len(mismatches)}):")
for m in mismatches: print("  ", m)
print(f"\nGaps — class sheet expects something, teacher sheets show nothing ({len(gaps)}):")
for g in gaps: print("  ", g)
print(f"\nExtra — teacher sheets have someone, class sheet blank ({len(extra)}):")
for e in extra: print("  ", e)
