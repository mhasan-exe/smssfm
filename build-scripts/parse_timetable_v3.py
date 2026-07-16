import openpyxl, json, sys

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else '/mnt/user-data/uploads/Working_Timetables_2026_-_27_-_Copy.xlsx'
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

DAYS = ['Monday','Tuesday','Wed','Thursday','Friday']
DAY_NORM = {'Monday':'Monday','Tuesday':'Tuesday','Wed':'Wednesday','Thursday':'Thursday','Friday':'Friday'}
SHEETS = ['Eng & SST teachers', 'urdu & isl teachers', 'sciences teachers', 'mathematics teachers']

def get_val(ws, merged, r, c):
    for mr in merged:
        if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
            return ws.cell(row=mr.min_row, column=mr.min_col).value
    return ws.cell(row=r, column=c).value

teachers = {}
teacher_meta = {}

for sheetname in SHEETS:
    ws = wb[sheetname]
    merged = list(ws.merged_cells.ranges)
    anchors = []
    for mr in merged:
        if mr.min_row == mr.max_row and (mr.max_col - mr.min_col) >= 4:
            v = ws.cell(row=mr.min_row, column=mr.min_col).value
            if v:
                anchors.append((mr.min_row, mr.min_col, mr.max_col, str(v).strip()))
    anchors.sort()
    clusters = {}
    for a in anchors:
        clusters.setdefault(a[1], []).append(a)
    for col, items in clusters.items():
        items.sort()
        for i, (row, mincol, maxcol, name) in enumerate(items):
            end_row = items[i+1][0] if i+1 < len(items) else ws.max_row+1
            teachers.setdefault(name, {})
            teacher_meta[name] = sheetname
            for r in range(row, end_row):
                dayval = get_val(ws, merged, r, mincol-1)
                if dayval and str(dayval).strip() in DAYS:
                    day = DAY_NORM[str(dayval).strip()]
                    slots = []
                    for c in range(mincol, mincol+8):
                        v = get_val(ws, merged, r, c)
                        slots.append(str(v).strip() if v else '')
                    if day in teachers[name]:
                        existing = teachers[name][day]
                        teachers[name][day] = [e if e else s for e,s in zip(existing, slots)]
                    else:
                        teachers[name][day] = slots

with open('teachers_raw_v3.json','w') as f:
    json.dump({'teachers':teachers,'meta':teacher_meta}, f, indent=1)

print("Total teachers parsed:", len(teachers))
for t in sorted(teachers):
    print(t, '|', teacher_meta[t], '| days:', list(teachers[t].keys()))
