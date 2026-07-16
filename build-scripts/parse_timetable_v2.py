import openpyxl, json, sys

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else 'Working_Timetables_2026_-_27_-_Copy.xlsx'
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb['Teachers Timetable']
merged = list(ws.merged_cells.ranges)

DAYS = ['Monday','Tuesday','Wed','Thursday','Friday']
DAY_NORM = {'Monday':'Monday','Tuesday':'Tuesday','Wed':'Wednesday','Thursday':'Thursday','Friday':'Friday'}

def get_val(r, c):
    for mr in merged:
        if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
            return ws.cell(row=mr.min_row, column=mr.min_col).value
    return ws.cell(row=r, column=c).value

anchors = []
for mr in merged:
    if mr.min_row == mr.max_row and (mr.max_col - mr.min_col) >= 4:
        v = ws.cell(row=mr.min_row, column=mr.min_col).value
        if v:
            anchors.append((mr.min_row, mr.min_col, mr.max_col, str(v).strip()))
anchors.sort()

clusters = {}
for a in anchors:
    clusters.setdefault(a[1], []).append(a)

teachers = {}
teacher_meta = {}
for col, items in clusters.items():
    items.sort()
    for i, (row, mincol, maxcol, name) in enumerate(items):
        end_row = items[i+1][0] if i+1 < len(items) else ws.max_row+1
        teachers.setdefault(name, {})
        teacher_meta[name] = 'Teachers Timetable'
        for r in range(row, end_row):
            dayval = get_val(r, mincol-1)
            if dayval and str(dayval).strip() in DAYS:
                day = DAY_NORM[str(dayval).strip()]
                slots = []
                for c in range(mincol, mincol+8):
                    v = get_val(r, c)
                    slots.append(str(v).strip() if v else '')
                if day in teachers[name]:
                    existing = teachers[name][day]
                    merged_slots = [e if e else s for e,s in zip(existing, slots)]
                    teachers[name][day] = merged_slots
                else:
                    teachers[name][day] = slots

with open('teachers_raw_v2.json','w') as f:
    json.dump({'teachers':teachers, 'meta':teacher_meta}, f, indent=1)

print("Total teachers parsed:", len(teachers))
for t in sorted(teachers):
    days_present = list(teachers[t].keys())
    print(t, '| days:', days_present)
